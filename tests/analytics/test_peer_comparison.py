"""Tests for Sprint 3 - Day 20 Peer Comparison Excel."""

import sqlite3

from openpyxl import load_workbook

from src.analytics.peer import PEER_METRICS
from src.analytics.peer_comparison import (
    METRIC_LABELS,
    build_peer_records,
    generate_peer_comparison_workbook,
    group_records,
    load_peer_comparison_data,
    safe_sheet_name,
)


# ============================================================
# SOURCE DATA
# ============================================================

def test_peer_comparison_data_loaded():
    rows = load_peer_comparison_data()

    assert len(rows) == 840


def test_source_has_11_peer_groups():
    rows = load_peer_comparison_data()

    groups = {
        row[0]
        for row in rows
    }

    assert len(groups) == 11


def test_source_has_56_companies():
    rows = load_peer_comparison_data()

    companies = {
        (row[0], row[1])
        for row in rows
    }

    assert len(companies) == 56


def test_source_has_15_metrics():
    rows = load_peer_comparison_data()

    metrics = {
        row[5]
        for row in rows
    }

    assert len(metrics) == 15
    assert metrics == set(
        PEER_METRICS
    )


# ============================================================
# COMPANY RECORDS
# ============================================================

def test_build_56_records():
    rows = load_peer_comparison_data()

    records = build_peer_records(
        rows
    )

    assert len(records) == 56


def test_each_record_has_15_metrics():
    rows = load_peer_comparison_data()

    records = build_peer_records(
        rows
    )

    for record in records:
        assert len(
            record["metrics"]
        ) == 15


def test_records_have_required_fields():
    rows = load_peer_comparison_data()

    records = build_peer_records(
        rows
    )

    record = records[0]

    assert {
        "peer_group_name",
        "company_id",
        "company_name",
        "is_benchmark",
        "year",
        "metrics",
    }.issubset(record)


# ============================================================
# GROUPING
# ============================================================

def test_group_records_has_11_groups():
    rows = load_peer_comparison_data()

    records = build_peer_records(
        rows
    )

    groups = group_records(
        records
    )

    assert len(groups) == 11


def test_group_company_counts():
    rows = load_peer_comparison_data()

    groups = group_records(
        build_peer_records(rows)
    )

    expected = {
        "Private Banks": 5,
        "Public Sector Banks": 4,
        "IT Services": 5,
        "Pharmaceuticals": 5,
        "Automobiles": 7,
        "Life Insurance": 4,
        "Oil & Gas": 5,
        "Power & Utilities": 7,
        "Steel": 4,
        "FMCG": 7,
        "Consumer Finance": 3,
    }

    assert {
        name: len(members)
        for name, members
        in groups.items()
    } == expected


def test_one_benchmark_per_group():
    rows = load_peer_comparison_data()

    groups = group_records(
        build_peer_records(rows)
    )

    for members in groups.values():

        benchmarks = [
            row
            for row in members
            if row["is_benchmark"]
        ]

        assert len(benchmarks) == 1


def test_expected_benchmarks():
    rows = load_peer_comparison_data()

    groups = group_records(
        build_peer_records(rows)
    )

    expected = {
        "Private Banks": "HDFCBANK",
        "Public Sector Banks": "SBIN",
        "IT Services": "TCS",
        "Pharmaceuticals": "SUNPHARMA",
        "Automobiles": "MARUTI",
        "Life Insurance": "LICI",
        "Oil & Gas": "RELIANCE",
        "Power & Utilities": "NTPC",
        "Steel": "TATASTEEL",
        "FMCG": "HINDUNILVR",
        "Consumer Finance": "BAJFINANCE",
    }

    for group_name, expected_id in (
        expected.items()
    ):

        actual = next(
            row["company_id"]
            for row
            in groups[group_name]
            if row["is_benchmark"]
        )

        assert actual == expected_id


# ============================================================
# METRIC CONFIGURATION
# ============================================================

def test_metric_count():
    assert len(PEER_METRICS) == 15


def test_all_metrics_have_labels():
    for metric in PEER_METRICS:
        assert metric in METRIC_LABELS


def test_metric_labels_not_empty():
    for metric in PEER_METRICS:
        assert METRIC_LABELS[
            metric
        ].strip()


# ============================================================
# SHEET NAME HANDLING
# ============================================================

def test_safe_sheet_name_normal():
    assert safe_sheet_name(
        "Private Banks"
    ) == "Private Banks"


def test_safe_sheet_name_invalid_chars():
    result = safe_sheet_name(
        "Bank/Test:*?"
    )

    assert result == (
        "Bank_Test___"
    )


def test_safe_sheet_name_max_31_chars():
    result = safe_sheet_name(
        "A" * 50
    )

    assert len(result) == 31


# ============================================================
# WORKBOOK GENERATION
# ============================================================

def test_workbook_generated(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    result = (
        generate_peer_comparison_workbook(
            output_path=path
        )
    )

    assert result["path"].exists()
    assert (
        result["path"].stat().st_size
        > 0
    )


def test_workbook_reports_11_groups(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    result = (
        generate_peer_comparison_workbook(
            output_path=path
        )
    )

    assert result[
        "peer_groups"
    ] == 11


def test_workbook_reports_56_companies(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    result = (
        generate_peer_comparison_workbook(
            output_path=path
        )
    )

    assert result[
        "companies"
    ] == 56


def test_workbook_reports_15_metrics(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    result = (
        generate_peer_comparison_workbook(
            output_path=path
        )
    )

    assert result[
        "metrics"
    ] == 15


def test_workbook_has_12_sheets(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    result = (
        generate_peer_comparison_workbook(
            output_path=path
        )
    )

    assert result[
        "sheets"
    ] == 12

    workbook = load_workbook(
        path,
        read_only=True,
    )

    assert len(
        workbook.sheetnames
    ) == 12

    workbook.close()


def test_summary_sheet_exists(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    assert (
        "Peer Group Summary"
        in workbook.sheetnames
    )

    workbook.close()


def test_all_peer_group_sheets_exist(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    expected = {
        "Automobiles",
        "Consumer Finance",
        "FMCG",
        "IT Services",
        "Life Insurance",
        "Oil & Gas",
        "Pharmaceuticals",
        "Power & Utilities",
        "Private Banks",
        "Public Sector Banks",
        "Steel",
    }

    actual = set(
        workbook.sheetnames
    )

    actual.remove(
        "Peer Group Summary"
    )

    assert actual == expected

    workbook.close()


# ============================================================
# SUMMARY SHEET STRUCTURE
# ============================================================

def test_summary_has_11_group_rows(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Peer Group Summary"
    ]

    # Header + 11 peer groups.
    assert sheet.max_row == 12

    workbook.close()


def test_summary_has_four_columns(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Peer Group Summary"
    ]

    assert sheet.max_column == 4

    workbook.close()


# ============================================================
# PEER SHEET STRUCTURE
# ============================================================

def test_peer_sheet_column_count(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "IT Services"
    ]

    # 4 company fields +
    # 15 metrics × 2 columns each.
    assert sheet.max_column == 34

    workbook.close()


def test_it_services_has_five_companies(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "IT Services"
    ]

    # Header + 5 companies.
    assert sheet.max_row == 6

    workbook.close()


def test_consumer_finance_has_three_companies(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Consumer Finance"
    ]

    # Header + 3 companies.
    assert sheet.max_row == 4

    workbook.close()


def test_benchmark_is_first_company(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "IT Services"
    ]

    assert sheet["A2"].value == "TCS"
    assert sheet["C2"].value == "YES"

    workbook.close()


# ============================================================
# RAW + PERCENTILE COLUMNS
# ============================================================

def test_first_metric_headers(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "IT Services"
    ]

    assert sheet["E1"].value == (
        "ROE (%)"
    )

    assert sheet["F1"].value == (
        "ROE (%) Percentile"
    )

    workbook.close()


def test_all_metric_header_pairs(
    tmp_path,
):
    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Private Banks"
    ]

    column = 5

    for metric in PEER_METRICS:

        label = METRIC_LABELS[
            metric
        ]

        assert sheet.cell(
            row=1,
            column=column,
        ).value == label

        assert sheet.cell(
            row=1,
            column=column + 1,
        ).value == (
            f"{label} Percentile"
        )

        column += 2

    workbook.close()


# ============================================================
# DATABASE INTEGRITY
# ============================================================

def get_peer_row_count():
    connection = sqlite3.connect(
        "nifty100.db"
    )

    try:
        return connection.execute(
            """
            SELECT COUNT(*)
            FROM peer_percentiles
            """
        ).fetchone()[0]

    finally:
        connection.close()


def test_peer_table_still_has_840_rows():
    assert get_peer_row_count() == 840


def test_excel_generation_does_not_modify_peer_table(
    tmp_path,
):
    before = get_peer_row_count()

    path = (
        tmp_path
        / "peer_comparison.xlsx"
    )

    generate_peer_comparison_workbook(
        output_path=path
    )

    after = get_peer_row_count()

    assert before == after == 840