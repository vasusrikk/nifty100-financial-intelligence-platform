"""Tests for Original Sprint 3 - Day 17 sector-relative ranking."""

from openpyxl import load_workbook

from src.screener.engine import load_screener_data
from src.screener.sector_ranking import (
    CATEGORY_METRICS,
    CATEGORY_WEIGHTS,
    LOWER_IS_BETTER,
    build_sector_bounds,
    group_by_sector,
    normalize_value,
    rank_sector_relative,
    score_company,
)
from src.screener.screener_excel import (
    EXPORT_COLUMNS,
    PRESETS,
    generate_screener_workbook,
)


# ============================================================
# 50 / 30 / 20 CONFIGURATION
# ============================================================

def test_category_weights():
    assert CATEGORY_WEIGHTS == {
        "profitability": 0.50,
        "growth": 0.30,
        "valuation": 0.20,
    }


def test_category_weights_total_one():
    assert round(
        sum(CATEGORY_WEIGHTS.values()),
        10,
    ) == 1.0


def test_category_metric_groups():
    assert set(CATEGORY_METRICS) == {
        "profitability",
        "growth",
        "valuation",
    }


def test_valuation_direction():
    assert "pe_ratio" in LOWER_IS_BETTER
    assert "pb_ratio" in LOWER_IS_BETTER


# ============================================================
# NORMALIZATION
# ============================================================

def test_higher_is_better_normalization():
    assert normalize_value(
        100,
        0,
        100,
    ) == 100.0

    assert normalize_value(
        0,
        0,
        100,
    ) == 0.0


def test_lower_is_better_normalization():
    assert normalize_value(
        0,
        0,
        100,
        lower_is_better=True,
    ) == 100.0

    assert normalize_value(
        100,
        0,
        100,
        lower_is_better=True,
    ) == 0.0


def test_equal_bounds_are_neutral():
    assert normalize_value(
        25,
        25,
        25,
    ) == 50.0


def test_missing_value_returns_none():
    assert normalize_value(
        None,
        0,
        100,
    ) is None


# ============================================================
# ACTUAL SECTOR DATA
# ============================================================

def test_current_broad_sector_count():
    rows = load_screener_data()

    groups = group_by_sector(rows)

    # Current database contains 10 broad sectors.
    assert len(groups) == 10


def test_all_companies_grouped():
    rows = load_screener_data()

    groups = group_by_sector(rows)

    assert sum(
        len(companies)
        for companies in groups.values()
    ) == 92


def test_financial_sector_count():
    rows = load_screener_data()

    groups = group_by_sector(rows)

    assert len(
        groups["Financials"]
    ) == 23


def test_sector_bounds_exist():
    rows = load_screener_data()

    bounds = build_sector_bounds(rows)

    assert "Financials" in bounds
    assert "roe_pct" in bounds["Financials"]
    assert "pe_ratio" in bounds["Financials"]


# ============================================================
# COMPANY SCORING
# ============================================================

def test_score_company_has_three_categories():
    rows = load_screener_data()
    bounds = build_sector_bounds(rows)

    scored = score_company(
        rows[0],
        bounds,
    )

    assert (
        "profitability_score"
        in scored
    )
    assert "growth_score" in scored
    assert "valuation_score" in scored


def test_score_company_has_composite():
    rows = load_screener_data()
    bounds = build_sector_bounds(rows)

    scored = score_company(
        rows[0],
        bounds,
    )

    assert (
        "composite_score_50_30_20"
        in scored
    )


def test_composite_score_range():
    ranked = rank_sector_relative()

    for row in ranked:
        assert (
            0
            <= row[
                "composite_score_50_30_20"
            ]
            <= 100
        )


# ============================================================
# COMPLETE RANKING
# ============================================================

def test_all_92_companies_ranked():
    ranked = rank_sector_relative()

    assert len(ranked) == 92


def test_overall_ranks_sequential():
    ranked = rank_sector_relative()

    assert [
        row["overall_rank"]
        for row in ranked
    ] == list(
        range(1, 93)
    )


def test_company_ids_unique():
    ranked = rank_sector_relative()

    ids = [
        row["company_id"]
        for row in ranked
    ]

    assert len(ids) == len(set(ids))


def test_ranking_descending():
    ranked = rank_sector_relative()

    scores = [
        row[
            "composite_score_50_30_20"
        ]
        for row in ranked
    ]

    assert scores == sorted(
        scores,
        reverse=True,
    )


def test_sector_rank_present():
    ranked = rank_sector_relative()

    for row in ranked:
        assert row["sector_rank"] >= 1


def test_each_sector_starts_at_rank_one():
    ranked = rank_sector_relative()

    sectors = {
        row["broad_sector"]
        for row in ranked
    }

    for sector in sectors:

        sector_rows = [
            row
            for row in ranked
            if row["broad_sector"]
            == sector
        ]

        assert min(
            row["sector_rank"]
            for row in sector_rows
        ) == 1


def test_ranking_deterministic():
    first = rank_sector_relative()
    second = rank_sector_relative()

    assert [
        row["company_id"]
        for row in first
    ] == [
        row["company_id"]
        for row in second
    ]


# ============================================================
# DAY 17 EXCEL CONFIGURATION
# ============================================================

def test_six_presets_configured():
    assert len(PRESETS) == 6


def test_export_has_more_than_20_fields():
    assert len(EXPORT_COLUMNS) >= 20


def test_required_score_columns_exported():
    fields = {
        field
        for field, _ in EXPORT_COLUMNS
    }

    assert (
        "profitability_score"
        in fields
    )
    assert "growth_score" in fields
    assert "valuation_score" in fields
    assert (
        "composite_score_50_30_20"
        in fields
    )


# ============================================================
# WORKBOOK GENERATION
# ============================================================

def test_workbook_generated(tmp_path):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    result = generate_screener_workbook(
        path
    )

    assert result["path"].exists()
    assert (
        result["path"].stat().st_size
        > 0
    )


def test_workbook_has_seven_sheets(
    tmp_path,
):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    result = generate_screener_workbook(
        path
    )

    assert result["sheet_count"] == 7


def test_workbook_sheet_names(
    tmp_path,
):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    generate_screener_workbook(path)

    workbook = load_workbook(
        path,
        read_only=True,
    )

    expected = [
        "Ranked Universe",
        "Quality Compounder",
        "Value Pick",
        "Growth Accelerator",
        "Dividend Champion",
        "Debt-Free Blue Chip",
        "Turnaround Watch",
    ]

    assert workbook.sheetnames == expected

    workbook.close()


def test_ranked_universe_has_92_rows(
    tmp_path,
):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    generate_screener_workbook(path)

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Ranked Universe"
    ]

    # 92 companies + header.
    assert sheet.max_row == 93

    workbook.close()


def test_workbook_column_count(
    tmp_path,
):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    generate_screener_workbook(path)

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Ranked Universe"
    ]

    assert sheet.max_column == len(
        EXPORT_COLUMNS
    )

    workbook.close()


def test_value_pick_sheet_has_two_companies(
    tmp_path,
):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    generate_screener_workbook(path)

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook["Value Pick"]

    # 2 companies + header.
    assert sheet.max_row == 3

    workbook.close()


def test_debt_free_sheet_has_two_companies(
    tmp_path,
):
    path = (
        tmp_path
        / "screener_output.xlsx"
    )

    generate_screener_workbook(path)

    workbook = load_workbook(
        path,
        read_only=True,
    )

    sheet = workbook[
        "Debt-Free Blue Chip"
    ]

    # Current corrected strict D/E == 0 result:
    # LICI and SBILIFE.
    assert sheet.max_row == 3

    workbook.close()