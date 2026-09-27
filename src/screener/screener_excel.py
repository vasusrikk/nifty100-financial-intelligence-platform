"""Sprint 3 - Original Day 17: Screener Excel Export.

Generates:
    reports/screener_output.xlsx

Workbook contains:
    - Ranked Universe
    - One worksheet for each of the 6 preset screeners
    - 50/30/20 sector-relative composite scores
    - 20+ financial / company fields
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.screener.engine import load_screener_data
from src.screener.presets import run_preset
from src.screener.sector_ranking import rank_sector_relative


# ============================================================
# WORKBOOK CONFIGURATION
# ============================================================

PRESETS = [
    ("quality_compounder", "Quality Compounder"),
    ("value_pick", "Value Pick"),
    ("growth_accelerator", "Growth Accelerator"),
    ("dividend_champion", "Dividend Champion"),
    ("debt_free_blue_chip", "Debt-Free Blue Chip"),
    ("turnaround_watch", "Turnaround Watch"),
]


EXPORT_COLUMNS = [
    ("overall_rank", "Overall Rank"),
    ("sector_rank", "Sector Rank"),
    ("company_id", "Company ID"),
    ("company_name", "Company Name"),
    ("year", "Financial Period"),
    ("broad_sector", "Broad Sector"),
    ("sub_sector", "Sub Sector"),

    ("profitability_score", "Profitability Score"),
    ("growth_score", "Growth Score"),
    ("valuation_score", "Valuation Score"),
    (
        "composite_score_50_30_20",
        "Composite Score",
    ),
    ("score_coverage_pct", "Score Coverage (%)"),

    ("roe_pct", "ROE (%)"),
    ("roce_pct", "ROCE (%)"),
    ("npm_pct", "Net Profit Margin (%)"),
    ("opm_pct", "Operating Profit Margin (%)"),
    ("de_ratio", "Debt / Equity"),
    ("icr", "Interest Coverage Ratio"),
    ("asset_turnover", "Asset Turnover"),

    ("revenue_cagr_5yr", "Revenue CAGR 5Y (%)"),
    ("pat_cagr_5yr", "PAT CAGR 5Y (%)"),
    ("eps_cagr_5yr", "EPS CAGR 5Y (%)"),

    ("cfo_margin_pct", "CFO Margin (%)"),
    ("cfo_pat_ratio", "CFO / PAT"),
    ("fcf", "Free Cash Flow"),
    ("fcf_margin_pct", "FCF Margin (%)"),
    ("capex_sales_pct", "Capex / Sales (%)"),

    ("pe_ratio", "P/E"),
    ("pb_ratio", "P/B"),
    ("dividend_yield_pct", "Dividend Yield (%)"),
    ("market_cap_crore", "Market Cap (₹ Cr)"),
]


# ============================================================
# LOOKUP HELPERS
# ============================================================

def build_rank_lookup(ranked_rows):
    """Create company_id -> ranked row mapping."""

    return {
        row["company_id"]: row
        for row in ranked_rows
    }


def merge_with_ranking(
    rows,
    rank_lookup,
):
    """
    Attach Day 17 sector-relative scoring fields to rows
    returned by preset screeners.
    """

    merged = []

    for row in rows:

        company_id = row.get(
            "company_id"
        )

        ranked = rank_lookup.get(
            company_id
        )

        if ranked is None:
            continue

        combined = dict(row)

        combined.update({
            "overall_rank":
                ranked.get("overall_rank"),

            "sector_rank":
                ranked.get("sector_rank"),

            "profitability_score":
                ranked.get(
                    "profitability_score"
                ),

            "growth_score":
                ranked.get(
                    "growth_score"
                ),

            "valuation_score":
                ranked.get(
                    "valuation_score"
                ),

            "composite_score_50_30_20":
                ranked.get(
                    "composite_score_50_30_20"
                ),

            "score_coverage_pct":
                ranked.get(
                    "score_coverage_pct"
                ),
        })

        merged.append(combined)

    merged.sort(
        key=lambda row: (
            row.get("overall_rank")
            if row.get("overall_rank")
            is not None
            else 999999
        )
    )

    return merged


# ============================================================
# WORKSHEET WRITER
# ============================================================

def write_sheet(
    workbook,
    title,
    rows,
):
    """Write and format one screener worksheet."""

    worksheet = workbook.create_sheet(
        title=title
    )

    headers = [
        label
        for _, label in EXPORT_COLUMNS
    ]

    worksheet.append(headers)

    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for row in rows:

        worksheet.append([
            row.get(field)
            for field, _ in EXPORT_COLUMNS
        ])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # --------------------------------------------------------
    # Alignment
    # --------------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top"
            )

    # --------------------------------------------------------
    # Number formatting
    # --------------------------------------------------------

    numeric_fields = {
        field
        for field, _ in EXPORT_COLUMNS
        if field not in {
            "company_id",
            "company_name",
            "year",
            "broad_sector",
            "sub_sector",
        }
    }

    for column_index, (
        field,
        _,
    ) in enumerate(
        EXPORT_COLUMNS,
        start=1,
    ):

        if field not in numeric_fields:
            continue

        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
            ).number_format = "0.00"

    # Rank columns should be integers.
    for column_index in (1, 2):
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
            ).number_format = "0"

    # --------------------------------------------------------
    # Composite score colour scale
    # --------------------------------------------------------

    composite_column = next(
        index
        for index, (field, _) in enumerate(
            EXPORT_COLUMNS,
            start=1,
        )
        if field
        == "composite_score_50_30_20"
    )

    if worksheet.max_row >= 2:

        column_letter = get_column_letter(
            composite_column
        )

        worksheet.conditional_formatting.add(
            (
                f"{column_letter}2:"
                f"{column_letter}"
                f"{worksheet.max_row}"
            ),
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    for column_index, (
        field,
        label,
    ) in enumerate(
        EXPORT_COLUMNS,
        start=1,
    ):

        if field == "company_name":
            width = 38
        elif field in {
            "broad_sector",
            "sub_sector",
        }:
            width = 25
        elif field == "company_id":
            width = 18
        else:
            width = max(
                13,
                min(
                    len(label) + 3,
                    24,
                ),
            )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width

    worksheet.row_dimensions[1].height = 35

    return worksheet


# ============================================================
# PRESET EXTRACTION
# ============================================================

def get_preset_rows(
    preset_name,
):
    """
    Return usable rows from one preset.

    SOURCE_LIMITATION presets still contain partial candidates,
    so those candidates are preserved in the workbook.
    """

    result = run_preset(
        preset_name
    )

    if result["status"] == "COMPLETE":
        rows = result.get(
            "results",
            []
        )
    else:
        rows = (
            result.get(
                "partial_results"
            )
            or result.get(
                "results"
            )
            or []
        )

    return result, rows


# ============================================================
# WORKBOOK GENERATION
# ============================================================

def generate_screener_workbook(
    output_path=(
        "reports/screener_output.xlsx"
    ),
):
    """
    Generate the original Sprint 3 Day 17 workbook.
    """

    universe = load_screener_data()

    ranked = rank_sector_relative(
        universe
    )

    rank_lookup = build_rank_lookup(
        ranked
    )

    output_path = Path(
        output_path
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    # Remove default blank worksheet.
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # --------------------------------------------------------
    # Sheet 1: complete ranked universe
    # --------------------------------------------------------

    write_sheet(
        workbook,
        "Ranked Universe",
        ranked,
    )

    preset_summary = []

    # --------------------------------------------------------
    # Six preset worksheets
    # --------------------------------------------------------

    for (
        preset_name,
        sheet_name,
    ) in PRESETS:

        result, rows = get_preset_rows(
            preset_name
        )

        merged = merge_with_ranking(
            rows,
            rank_lookup,
        )

        write_sheet(
            workbook,
            sheet_name,
            merged,
        )

        preset_summary.append({
            "preset": sheet_name,
            "status": result.get(
                "status"
            ),
            "companies": len(merged),
            "reason": result.get(
                "reason"
            ),
        })

    workbook.save(
        output_path
    )

    return {
        "path": output_path,
        "universe_count": len(
            universe
        ),
        "ranked_count": len(
            ranked
        ),
        "preset_summary":
            preset_summary,
        "sheet_count":
            len(workbook.sheetnames),
        "sheet_names":
            workbook.sheetnames,
    }


# ============================================================
# DAY 17 QA
# ============================================================

def main():

    result = (
        generate_screener_workbook()
    )

    path = result["path"]

    print(
        "\n=== ORIGINAL SPRINT 3 - "
        "DAY 17 EXCEL QA ==="
    )

    print(
        "Universe:",
        result["universe_count"],
    )

    print(
        "Ranked:",
        result["ranked_count"],
    )

    print(
        "Sheets:",
        result["sheet_count"],
    )

    print(
        "Workbook:",
        path,
    )

    print(
        "Workbook exists:",
        path.exists(),
    )

    print(
        "Workbook size:",
        path.stat().st_size,
        "bytes",
    )

    print(
        "\nSHEETS"
    )

    for sheet_name in (
        result["sheet_names"]
    ):
        print(
            "-",
            sheet_name,
        )

    print(
        "\nPRESET SUMMARY"
    )

    for preset in (
        result["preset_summary"]
    ):

        print(
            preset["preset"],
            "|",
            preset["status"],
            "| Companies:",
            preset["companies"],
        )

        if preset["reason"]:
            print(
                "  Reason:",
                preset["reason"],
            )


if __name__ == "__main__":
    main()