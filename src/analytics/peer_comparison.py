"""Sprint 3 - Day 20: Peer Comparison Excel Report."""

from __future__ import annotations

import sqlite3
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.analytics.peer import PEER_METRICS


# ============================================================
# CONFIGURATION
# ============================================================

DATABASE_PATH = Path("nifty100.db")

OUTPUT_PATH = Path(
    "reports/peer_comparison.xlsx"
)


METRIC_LABELS = {
    "roe_pct": "ROE (%)",
    "roce_pct": "ROCE (%)",
    "npm_pct": "Net Profit Margin (%)",
    "opm_pct": "Operating Profit Margin (%)",
    "asset_turnover": "Asset Turnover",
    "revenue_cagr_5yr": "Revenue CAGR 5Y (%)",
    "pat_cagr_5yr": "PAT CAGR 5Y (%)",
    "eps_cagr_5yr": "EPS CAGR 5Y (%)",
    "cfo_margin_pct": "CFO Margin (%)",
    "icr": "Interest Coverage Ratio",
    "de_ratio": "Debt / Equity",
    "pe_ratio": "P/E",
    "pb_ratio": "P/B",
    "dividend_yield_pct": "Dividend Yield (%)",
    "market_cap_crore": "Market Cap (₹ Cr)",
}


# ============================================================
# LOAD PEER DATA
# ============================================================

def load_peer_comparison_data(
    database_path=DATABASE_PATH,
):
    """
    Load Day 18 peer percentile data together with
    company names.

    This function is read-only.
    """

    connection = sqlite3.connect(
        database_path
    )

    try:
        rows = connection.execute(
            """
            SELECT
                p.peer_group_name,
                p.company_id,
                c.company_name,
                p.is_benchmark,
                p.year,
                p.metric,
                p.raw_value,
                p.percentile

            FROM peer_percentiles AS p

            LEFT JOIN companies AS c
                ON c.id = p.company_id

            ORDER BY
                p.peer_group_name,
                p.company_id,
                p.metric
            """
        ).fetchall()

    finally:
        connection.close()

    return rows


# ============================================================
# BUILD COMPANY RECORDS
# ============================================================

def build_peer_records(rows):
    """
    Convert normalized percentile rows into one company
    record per peer-group membership.
    """

    records = {}

    for (
        peer_group,
        company_id,
        company_name,
        is_benchmark,
        year,
        metric,
        raw_value,
        percentile,
    ) in rows:

        key = (
            peer_group,
            company_id,
        )

        if key not in records:
            records[key] = {
                "peer_group_name":
                    peer_group,

                "company_id":
                    company_id,

                "company_name":
                    company_name
                    or company_id,

                "is_benchmark":
                    bool(is_benchmark),

                "year":
                    year,

                "metrics": {},
            }

        records[key]["metrics"][metric] = {
            "raw_value":
                raw_value,

            "percentile":
                percentile,
        }

    return list(
        records.values()
    )


def group_records(records):
    """
    Group company records by peer group.
    """

    groups = defaultdict(list)

    for record in records:
        groups[
            record["peer_group_name"]
        ].append(record)

    return dict(groups)


# ============================================================
# EXCEL HELPERS
# ============================================================

def safe_sheet_name(name):
    """
    Convert a peer-group name into a valid Excel
    worksheet name.
    """

    invalid = set('[]:*?/\\')

    cleaned = "".join(
        "_"
        if character in invalid
        else character
        for character in str(name)
    )

    return cleaned[:31]


def style_header(row):
    """
    Apply formatting to worksheet header cells.
    """

    fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in row:

        cell.fill = fill
        cell.font = font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def autosize_columns(
    worksheet,
    max_width=35,
):
    """
    Resize worksheet columns for readability.
    """

    for column_cells in worksheet.columns:

        maximum_length = 0

        for cell in column_cells:

            value = cell.value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        column_letter = (
            get_column_letter(
                column_cells[0].column
            )
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 2,
            max_width,
        )


# ============================================================
# SUMMARY SHEET
# ============================================================

def add_summary_sheet(
    workbook,
    groups,
):
    """
    Add the peer-group summary worksheet.
    """

    worksheet = workbook.active

    worksheet.title = (
        "Peer Group Summary"
    )

    worksheet.append([
        "Peer Group",
        "Companies",
        "Benchmark",
        "Metrics",
    ])

    style_header(
        worksheet[1]
    )

    for group_name in sorted(groups):

        members = groups[
            group_name
        ]

        benchmark = next(
            (
                member["company_id"]
                for member in members
                if member["is_benchmark"]
            ),
            None,
        )

        worksheet.append([
            group_name,
            len(members),
            benchmark,
            len(PEER_METRICS),
        ])

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    autosize_columns(
        worksheet
    )


# ============================================================
# INDIVIDUAL PEER GROUP SHEET
# ============================================================

def add_peer_group_sheet(
    workbook,
    group_name,
    records,
):
    """
    Create one worksheet for one peer group.

    Each financial metric contains:
    1. Raw financial value
    2. Peer percentile
    """

    worksheet = (
        workbook.create_sheet(
            safe_sheet_name(
                group_name
            )
        )
    )

    headers = [
        "Company ID",
        "Company Name",
        "Benchmark",
        "Financial Year",
    ]

    for metric in PEER_METRICS:

        label = METRIC_LABELS.get(
            metric,
            metric,
        )

        headers.extend([
            label,
            f"{label} Percentile",
        ])

    worksheet.append(
        headers
    )

    style_header(
        worksheet[1]
    )

    # Benchmark appears first.
    # Remaining companies are sorted alphabetically.
    sorted_records = sorted(
        records,
        key=lambda row: (
            not row["is_benchmark"],
            row["company_id"],
        ),
    )

    for record in sorted_records:

        output = [
            record["company_id"],
            record["company_name"],
            (
                "YES"
                if record["is_benchmark"]
                else ""
            ),
            record["year"],
        ]

        for metric in PEER_METRICS:

            metric_data = (
                record["metrics"].get(
                    metric,
                    {},
                )
            )

            output.extend([
                metric_data.get(
                    "raw_value"
                ),
                metric_data.get(
                    "percentile"
                ),
            ])

        worksheet.append(
            output
        )

    worksheet.freeze_panes = "E2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Every second metric column is a percentile column.
    # Columns:
    # A-D = company information
    # E = first raw value
    # F = first percentile
    # G = second raw value
    # H = second percentile
    # etc.
    for column in range(
        6,
        worksheet.max_column + 1,
        2,
    ):

        for row in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row,
                column=column,
            ).number_format = "0.00"

    autosize_columns(
        worksheet
    )


# ============================================================
# GENERATE WORKBOOK
# ============================================================

def generate_peer_comparison_workbook(
    database_path=DATABASE_PATH,
    output_path=OUTPUT_PATH,
):
    """
    Generate the complete Sprint 3 Day 20
    peer-comparison Excel workbook.
    """

    rows = load_peer_comparison_data(
        database_path
    )

    if not rows:
        raise ValueError(
            "No peer percentile data found. "
            "Run Day 18 peer analytics first."
        )

    records = build_peer_records(
        rows
    )

    groups = group_records(
        records
    )

    if len(groups) != 11:
        raise ValueError(
            "Expected 11 peer groups, "
            f"found {len(groups)}."
        )

    if len(records) != 56:
        raise ValueError(
            "Expected 56 peer-group company "
            f"memberships, found {len(records)}."
        )

    output_path = Path(
        output_path
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    add_summary_sheet(
        workbook,
        groups,
    )

    for group_name in sorted(groups):

        add_peer_group_sheet(
            workbook,
            group_name,
            groups[group_name],
        )

    workbook.save(
        output_path
    )

    workbook.close()

    return {
        "path":
            output_path,

        "peer_groups":
            len(groups),

        "companies":
            len(records),

        "metrics":
            len(PEER_METRICS),

        "sheets":
            1 + len(groups),

        "records":
            records,
    }


# ============================================================
# DAY 20 QA
# ============================================================

def main():

    result = (
        generate_peer_comparison_workbook()
    )

    print(
        "\n=== ORIGINAL SPRINT 3 - "
        "DAY 20 PEER COMPARISON EXCEL QA ==="
    )

    print(
        "Peer groups:",
        result["peer_groups"],
    )

    print(
        "Companies:",
        result["companies"],
    )

    print(
        "Metrics:",
        result["metrics"],
    )

    print(
        "Sheets:",
        result["sheets"],
    )

    print(
        "Workbook:",
        result["path"],
    )

    print(
        "Workbook exists:",
        result["path"].exists(),
    )

    if result["path"].exists():

        print(
            "Workbook size:",
            result["path"].stat().st_size,
            "bytes",
        )

    print(
        "\nPEER GROUPS"
    )

    groups = group_records(
        result["records"]
    )

    for group_name in sorted(groups):

        members = groups[
            group_name
        ]

        benchmark = next(
            member["company_id"]
            for member in members
            if member["is_benchmark"]
        )

        print(
            group_name,
            "| Companies:",
            len(members),
            "| Benchmark:",
            benchmark,
        )


if __name__ == "__main__":
    main()