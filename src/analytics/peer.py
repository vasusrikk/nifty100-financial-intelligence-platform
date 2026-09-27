"""Sprint 3 - Day 18: Peer Group Percentile Analytics."""

from __future__ import annotations

import math
import sqlite3
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

DATABASE_PATH = Path("nifty100.db")

PEER_GROUPS_PATH = Path(
    "data/Supporting/"
    "1788501620796-5060f580-peer_groups.xlsx"
)

EXPECTED_PEER_GROUP_COUNT = 11


# Metrics used for peer percentile analysis.
# Higher percentile always means "better".
PEER_METRICS = {
    "roe_pct": True,
    "roce_pct": True,
    "npm_pct": True,
    "opm_pct": True,
    "asset_turnover": True,
    "revenue_cagr_5yr": True,
    "pat_cagr_5yr": True,
    "eps_cagr_5yr": True,
    "cfo_margin_pct": True,
    "icr": True,

    # Lower is better.
    "de_ratio": False,
    "pe_ratio": False,
    "pb_ratio": False,

    # Higher is better.
    "dividend_yield_pct": True,
    "market_cap_crore": True,
}


# ============================================================
# VALUE HELPERS
# ============================================================

def valid_number(value):
    """Return True only for finite numeric values."""

    if value is None:
        return False

    try:
        number = float(value)
    except (TypeError, ValueError):
        return False

    return math.isfinite(number)


# ============================================================
# PEER GROUP LOADER
# ============================================================

def load_peer_groups(
    path=PEER_GROUPS_PATH,
):
    """
    Load peer-group membership from the supplied workbook.
    """

    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            f"Peer group workbook not found: {path}"
        )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    expected_headers = [
        "id",
        "peer_group_name",
        "company_id",
        "is_benchmark",
    ]

    if headers != expected_headers:
        workbook.close()

        raise ValueError(
            "Unexpected peer_groups.xlsx columns. "
            f"Expected {expected_headers}, got {headers}"
        )

    rows = []

    for values in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        (
            row_id,
            peer_group_name,
            company_id,
            is_benchmark,
        ) = values

        if (
            peer_group_name is None
            or company_id is None
        ):
            continue

        rows.append({
            "id": row_id,
            "peer_group_name":
                str(peer_group_name).strip(),
            "company_id":
                str(company_id).strip(),
            "is_benchmark":
                bool(is_benchmark),
        })

    workbook.close()

    return rows


# ============================================================
# PEER GROUP VALIDATION
# ============================================================

def group_peer_memberships(
    memberships,
):
    """Group workbook rows by peer_group_name."""

    groups = defaultdict(list)

    for row in memberships:
        groups[
            row["peer_group_name"]
        ].append(row)

    return dict(groups)


def validate_peer_groups(
    memberships,
    expected_group_count=(
        EXPECTED_PEER_GROUP_COUNT
    ),
):
    """
    Validate the supplied peer-group mapping.

    Requirements:
    - exactly 11 groups
    - no duplicate company inside a group
    - exactly one benchmark per group
    """

    groups = group_peer_memberships(
        memberships
    )

    if len(groups) != expected_group_count:
        raise ValueError(
            "Expected "
            f"{expected_group_count} peer groups, "
            f"found {len(groups)}."
        )

    for group_name, members in (
        groups.items()
    ):

        company_ids = [
            member["company_id"]
            for member in members
        ]

        if len(company_ids) != len(
            set(company_ids)
        ):
            raise ValueError(
                "Duplicate company membership in "
                f"peer group: {group_name}"
            )

        benchmarks = [
            member
            for member in members
            if member["is_benchmark"]
        ]

        if len(benchmarks) != 1:
            raise ValueError(
                f"Peer group '{group_name}' "
                "must contain exactly one benchmark."
            )

    return groups


# ============================================================
# DATABASE VALIDATION
# ============================================================

def validate_company_ids(
    connection,
    memberships,
):
    """
    Ensure every peer workbook company exists in companies.id.
    """

    database_ids = {
        row[0]
        for row in connection.execute(
            "SELECT id FROM companies"
        ).fetchall()
    }

    missing = sorted({
        row["company_id"]
        for row in memberships
        if row["company_id"]
        not in database_ids
    })

    if missing:
        raise ValueError(
            "Peer group companies missing from "
            f"database: {missing}"
        )

    return True


# ============================================================
# LATEST FINANCIAL DATA
# ============================================================

def load_latest_peer_financials(
    connection,
    memberships,
):
    """
    Load the latest usable financial-ratio row for each
    peer-group company.

    Market-cap metrics are joined using the same financial
    period where available.
    """

    company_ids = sorted({
        row["company_id"]
        for row in memberships
    })

    placeholders = ",".join(
        "?"
        for _ in company_ids
    )

    query = f"""
        SELECT
            f.company_id,
            f.year,

            f.roe_pct,
            f.roce_pct,
            f.npm_pct,
            f.opm_pct,
            f.asset_turnover,

            f.revenue_cagr_5yr,
            f.pat_cagr_5yr,
            f.eps_cagr_5yr,

            f.cfo_margin_pct,
            f.icr,
            f.de_ratio,

            m.pe_ratio,
            m.pb_ratio,
            m.dividend_yield_pct,
            m.market_cap_crore

        FROM financial_ratios AS f

        LEFT JOIN market_cap AS m
            ON m.company_id = f.company_id
            AND m.year = f.year

        WHERE f.company_id IN (
            {placeholders}
        )

        AND f.year = (
            SELECT MAX(f2.year)
            FROM financial_ratios AS f2
            WHERE
                f2.company_id = f.company_id
                AND (
                    f2.roe_pct IS NOT NULL
                    OR f2.roce_pct IS NOT NULL
                    OR f2.npm_pct IS NOT NULL
                    OR f2.opm_pct IS NOT NULL
                    OR f2.revenue_cagr_5yr IS NOT NULL
                )
        )
    """

    cursor = connection.execute(
        query,
        company_ids,
    )

    columns = [
        description[0]
        for description in cursor.description
    ]

    rows = [
        dict(zip(columns, values))
        for values in cursor.fetchall()
    ]

    return {
        row["company_id"]: row
        for row in rows
    }


# ============================================================
# PERCENT_RANK
# ============================================================

def percent_rank(
    values,
    higher_is_better=True,
):
    """
    Calculate SQL-style PERCENT_RANK for a mapping:

        company_id -> numeric value

    Formula:
        (rank - 1) / (N - 1)

    Equal values receive the same rank.

    Returned percentile is on a 0-100 scale.

    For lower-is-better metrics, ordering is reversed so
    100 always represents the stronger peer-relative result.
    """

    clean = {
        company_id: float(value)
        for company_id, value
        in values.items()
        if valid_number(value)
    }

    if not clean:
        return {}

    if len(clean) == 1:
        company_id = next(
            iter(clean)
        )

        return {
            company_id: 100.0
        }

    ordered_values = sorted(
        set(clean.values()),
        reverse=higher_is_better,
    )

    rank_by_value = {}

    previous_count = 0

    for value in ordered_values:

        # SQL RANK behavior:
        # rank = 1 + number of rows preceding this value.
        rank_by_value[value] = (
            previous_count + 1
        )

        previous_count += sum(
            1
            for current
            in clean.values()
            if current == value
        )

    denominator = len(clean) - 1

    result = {}

    for company_id, value in (
        clean.items()
    ):

        rank = rank_by_value[
            value
        ]

        sql_percent_rank = (
            (rank - 1)
            / denominator
        )

        # SQL PERCENT_RANK gives the best ordered value 0.
        # Convert to "higher percentile = better".
        quality_percentile = (
            1.0 - sql_percent_rank
        ) * 100.0

        result[company_id] = round(
            quality_percentile,
            4,
        )

    return result


# ============================================================
# CALCULATE PEER PERCENTILES
# ============================================================

def calculate_peer_percentiles(
    memberships,
    financials,
):
    """
    Calculate all metric percentiles within each peer group.
    """

    groups = validate_peer_groups(
        memberships
    )

    output = []

    for group_name, members in (
        groups.items()
    ):

        member_ids = [
            member["company_id"]
            for member in members
        ]

        benchmark_id = next(
            member["company_id"]
            for member in members
            if member["is_benchmark"]
        )

        metric_percentiles = {}

        for (
            metric,
            higher_is_better,
        ) in PEER_METRICS.items():

            values = {
                company_id:
                    financials.get(
                        company_id,
                        {},
                    ).get(metric)

                for company_id
                in member_ids
            }

            metric_percentiles[
                metric
            ] = percent_rank(
                values,
                higher_is_better=(
                    higher_is_better
                ),
            )

        for company_id in member_ids:

            financial_row = (
                financials.get(
                    company_id
                )
            )

            if financial_row is None:
                year = None
            else:
                year = financial_row.get(
                    "year"
                )

            for metric in PEER_METRICS:

                raw_value = None

                if financial_row:
                    raw_value = (
                        financial_row.get(
                            metric
                        )
                    )

                percentile_value = (
                    metric_percentiles[
                        metric
                    ].get(company_id)
                )

                output.append({
                    "peer_group_name":
                        group_name,

                    "company_id":
                        company_id,

                    "is_benchmark":
                        company_id
                        == benchmark_id,

                    "year":
                        year,

                    "metric":
                        metric,

                    "raw_value":
                        raw_value,

                    "percentile":
                        percentile_value,
                })

    return output


# ============================================================
# DATABASE TABLE
# ============================================================

def create_peer_percentiles_table(
    connection,
):
    """
    Create the derived Day 18 percentile table.
    """

    connection.execute("""
        CREATE TABLE IF NOT EXISTS peer_percentiles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,

            peer_group_name TEXT NOT NULL,
            company_id TEXT NOT NULL,
            is_benchmark INTEGER NOT NULL DEFAULT 0,

            year TEXT,
            metric TEXT NOT NULL,

            raw_value REAL,
            percentile REAL,

            UNIQUE (
                peer_group_name,
                company_id,
                year,
                metric
            )
        )
    """)

    connection.execute("""
        CREATE INDEX IF NOT EXISTS
        idx_peer_percentiles_group
        ON peer_percentiles (
            peer_group_name
        )
    """)

    connection.execute("""
        CREATE INDEX IF NOT EXISTS
        idx_peer_percentiles_company
        ON peer_percentiles (
            company_id
        )
    """)


def save_peer_percentiles(
    connection,
    rows,
):
    """
    Replace the derived percentile table contents atomically.
    """

    create_peer_percentiles_table(
        connection
    )

    connection.execute(
        "DELETE FROM peer_percentiles"
    )

    connection.executemany("""
        INSERT INTO peer_percentiles (
            peer_group_name,
            company_id,
            is_benchmark,
            year,
            metric,
            raw_value,
            percentile
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, [
        (
            row["peer_group_name"],
            row["company_id"],
            int(row["is_benchmark"]),
            row["year"],
            row["metric"],
            row["raw_value"],
            row["percentile"],
        )
        for row in rows
    ])


# ============================================================
# DAY 18 PIPELINE
# ============================================================

def build_peer_percentiles(
    database_path=DATABASE_PATH,
    peer_groups_path=PEER_GROUPS_PATH,
):
    """
    Execute the complete Day 18 peer analytics pipeline.
    """

    memberships = load_peer_groups(
        peer_groups_path
    )

    groups = validate_peer_groups(
        memberships
    )

    connection = sqlite3.connect(
        database_path
    )

    try:
        validate_company_ids(
            connection,
            memberships,
        )

        financials = (
            load_latest_peer_financials(
                connection,
                memberships,
            )
        )

        rows = (
            calculate_peer_percentiles(
                memberships,
                financials,
            )
        )

        save_peer_percentiles(
            connection,
            rows,
        )

        connection.commit()

    except Exception:
        connection.rollback()
        raise

    finally:
        connection.close()

    return {
        "peer_groups": len(groups),
        "memberships": len(
            memberships
        ),
        "companies_with_financials":
            len(financials),
        "metrics": len(
            PEER_METRICS
        ),
        "percentile_rows": len(
            rows
        ),
        "rows": rows,
    }


# ============================================================
# QA
# ============================================================

def main():

    result = build_peer_percentiles()

    print(
        "\n=== ORIGINAL SPRINT 3 - "
        "DAY 18 PEER PERCENTILE QA ==="
    )

    print(
        "Peer groups:",
        result["peer_groups"],
    )

    print(
        "Memberships:",
        result["memberships"],
    )

    print(
        "Companies with financials:",
        result[
            "companies_with_financials"
        ],
    )

    print(
        "Metrics:",
        result["metrics"],
    )

    print(
        "Percentile rows:",
        result[
            "percentile_rows"
        ],
    )

    connection = sqlite3.connect(
        DATABASE_PATH
    )

    try:
        stored = connection.execute(
            """
            SELECT COUNT(*)
            FROM peer_percentiles
            """
        ).fetchone()[0]

        populated = connection.execute(
            """
            SELECT COUNT(*)
            FROM peer_percentiles
            WHERE percentile IS NOT NULL
            """
        ).fetchone()[0]

        print(
            "Rows stored:",
            stored,
        )

        print(
            "Populated percentiles:",
            populated,
        )

        print(
            "\nPEER GROUP SUMMARY"
        )

        summary = connection.execute("""
            SELECT
                peer_group_name,
                COUNT(DISTINCT company_id),
                SUM(
                    CASE
                    WHEN is_benchmark = 1
                    THEN 1
                    ELSE 0
                    END
                ) / COUNT(DISTINCT metric)
            FROM peer_percentiles
            GROUP BY peer_group_name
            ORDER BY peer_group_name
        """).fetchall()

        for (
            group_name,
            company_count,
            benchmark_count,
        ) in summary:

            print(
                group_name,
                "| Companies:",
                company_count,
                "| Benchmarks:",
                benchmark_count,
            )

    finally:
        connection.close()


if __name__ == "__main__":
    main()